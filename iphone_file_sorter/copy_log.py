"""Excel logging for iPhone copy results."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any


FAILED_HEADERS = [
    "Timestamp",
    "Device",
    "Source folder",
    "File name",
    "Relative path",
    "Source path",
    "Destination path",
    "Error",
    "Status",
]

SUCCESS_HEADERS = [
    "Timestamp",
    "Device",
    "Source folder",
    "File name",
    "Relative path",
    "Destination path",
    "Status",
]


def default_log_path(destination: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return Path(destination) / f"iphone_copy_log_{stamp}.xlsx"


def write_copy_excel_log(
    log_path: Path,
    *,
    device_name: str,
    selected_folders: list[str],
    succeeded: list[dict[str, Any]],
    failed: list[dict[str, Any]],
) -> Path:
    """
    Write an Excel workbook with:
    - Failed copies (details)
    - Successful copies
    - Summary
    """
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for Excel logs. Run: python -m pip install openpyxl"
        ) from exc

    log_path = Path(log_path)
    log_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    ws_failed = wb.active
    ws_failed.title = "Failed copies"
    ws_failed.append(FAILED_HEADERS)
    for row in failed:
        ws_failed.append(
            [
                row.get("timestamp", ""),
                row.get("device", device_name),
                row.get("source_folder", ""),
                row.get("file_name", ""),
                row.get("relative_path", ""),
                row.get("source_path", ""),
                row.get("destination_path", ""),
                row.get("error", ""),
                row.get("status", "Failed"),
            ]
        )

    ws_ok = wb.create_sheet("Successful copies")
    ws_ok.append(SUCCESS_HEADERS)
    for row in succeeded:
        ws_ok.append(
            [
                row.get("timestamp", ""),
                row.get("device", device_name),
                row.get("source_folder", ""),
                row.get("file_name", ""),
                row.get("relative_path", ""),
                row.get("destination_path", ""),
                row.get("status", "Copied"),
            ]
        )

    ws_sum = wb.create_sheet("Summary")
    ws_sum.append(["Field", "Value"])
    ws_sum.append(["Generated at", datetime.now().isoformat(timespec="seconds")])
    ws_sum.append(["Device", device_name])
    ws_sum.append(["Folders in this batch", ", ".join(selected_folders)])
    ws_sum.append(["Files copied successfully", len(succeeded)])
    ws_sum.append(["Files failed / skipped", len(failed)])
    ws_sum.append(
        [
            "Total files attempted",
            len(succeeded) + len(failed),
        ]
    )

    wb.save(log_path)
    return log_path
