#!/usr/bin/env python3
"""Tests for Excel copy logging."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from copy_log import write_copy_excel_log


class CopyLogTests(unittest.TestCase):
    def test_writes_failed_and_summary_sheets(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl not installed")

        with tempfile.TemporaryDirectory() as tmp:
            log_path = Path(tmp) / "log.xlsx"
            write_copy_excel_log(
                log_path,
                device_name="Test iPhone",
                selected_folders=["202403_b", "202403_c"],
                succeeded=[
                    {
                        "timestamp": "2026-07-26T10:00:00",
                        "device": "Test iPhone",
                        "source_folder": "202403_b",
                        "file_name": "a.jpg",
                        "relative_path": "a.jpg",
                        "destination_path": r"C:\out\202403_b\a.jpg",
                        "status": "Copied",
                    }
                ],
                failed=[
                    {
                        "timestamp": "2026-07-26T10:00:01",
                        "device": "Test iPhone",
                        "source_folder": "202403_b",
                        "file_name": "b.heic",
                        "relative_path": "b.heic",
                        "source_path": "Test iPhone/Internal Storage/202403_b/b.heic",
                        "destination_path": r"C:\out\202403_b\b.heic",
                        "error": "Timed out — file did not appear at destination",
                        "status": "Skipped",
                    }
                ],
            )

            self.assertTrue(log_path.exists())
            from openpyxl import load_workbook

            wb = load_workbook(log_path)
            self.assertIn("Failed copies", wb.sheetnames)
            self.assertIn("Successful copies", wb.sheetnames)
            self.assertIn("Summary", wb.sheetnames)
            failed = wb["Failed copies"]
            self.assertEqual(failed["A1"].value, "Timestamp")
            self.assertEqual(failed["D2"].value, "b.heic")
            self.assertEqual(failed["H2"].value, "Timed out — file did not appear at destination")


if __name__ == "__main__":
    unittest.main()
